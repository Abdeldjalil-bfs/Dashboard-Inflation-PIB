"""
Saisie de nouvelles valeurs brutes dans le fichier source.

Module sans dépendance à Streamlit, testable en isolation. Il ne calcule
aucun indicateur : après écriture, il délègue le recalcul complet à
backend.inflation.calculator.pipeline_global(), seule autorité en matière
d'IPC, d'inflation et de contributions.

La structure des paniers vient de config/categories.json — la même source
que visualizer.py — pour qu'un ajout d'élément ne puisse jamais désynchroniser
la saisie et l'affichage.
"""

import json
import warnings

import pandas as pd
from openpyxl import load_workbook

from config.settings import CATEGORIES_PATH, FICHIER_DONNEES, PANIERS
from backend.inflation.calculator import get_max_date, pipeline_global


class DateDejaPresente(ValueError):
    """Une ligne existe déjà pour ce mois et l'écrasement n'a pas été demandé."""


def _fichier(nom_fichier=None) -> str:
    return str(nom_fichier or FICHIER_DONNEES)


# ---------------------------------------------------------------------------
# Structure
# ---------------------------------------------------------------------------


def structure_panier(panier: str) -> list:
    """
    Éléments attendus pour ce panier, dans l'ordre de config/categories.json.
    """
    with open(CATEGORIES_PATH, "r", encoding="utf-8") as flux:
        categories = json.load(flux)
    if panier not in categories:
        raise ValueError("Panier inconnu dans categories.json : " + str(panier))
    return list(categories[panier].keys())


def paniers_disponibles(nom_fichier=None) -> list:
    """
    Paniers réellement présents dans le classeur, dans l'ordre de PANIERS.

    Le fichier brut ne porte pas la feuille 'core' : proposer sa saisie
    produirait une erreur à l'écriture.
    """
    feuilles = set(pd.ExcelFile(_fichier(nom_fichier)).sheet_names)
    return [p for p in PANIERS if p in feuilles]


def paniers_manquants(nom_fichier=None) -> list:
    """Paniers déclarés mais absents du classeur."""
    presents = set(paniers_disponibles(nom_fichier))
    return [p for p in PANIERS if p not in presents]


# ---------------------------------------------------------------------------
# Dates
# ---------------------------------------------------------------------------


def date_max_globale(nom_fichier=None) -> pd.Timestamp:
    """
    Dernier mois couvert par TOUS les paniers.

    On retient volontairement la PLUS ANCIENNE des dates maximales : si un
    panier est en retard, proposer une insertion au-delà de son dernier point
    creuserait un trou dans sa série.
    """
    chemin = _fichier(nom_fichier)
    maxima = {}
    for panier in paniers_disponibles(chemin):
        try:
            maxima[panier] = pd.Timestamp(get_max_date(chemin, panier))
        except Exception:
            continue

    if not maxima:
        raise ValueError("Aucun panier lisible dans " + chemin)

    plus_ancienne = min(maxima.values())
    plus_recente = max(maxima.values())
    if plus_ancienne != plus_recente:
        retard = [p for p, d in maxima.items() if d < plus_recente]
        warnings.warn(
            "Paniers désynchronisés : %s s'arrêtent à %s alors que d'autres "
            "vont jusqu'à %s. La saisie repart de la date la plus ancienne."
            % (", ".join(sorted(retard)), plus_ancienne.date(), plus_recente.date()),
            stacklevel=2,
        )
    return plus_ancienne


def proposer_dates_insertion(nom_fichier=None, nb_mois: int = 6) -> list:
    """Les `nb_mois` mois qui suivent le dernier mois couvert."""
    depart = pd.Period(date_max_globale(nom_fichier), freq="M")
    return [depart + rang for rang in range(1, int(nb_mois) + 1)]


# ---------------------------------------------------------------------------
# Lecture
# ---------------------------------------------------------------------------


def _lire_feuille(nom_fichier, panier) -> pd.DataFrame:
    df = pd.read_excel(_fichier(nom_fichier), sheet_name=panier)
    df.rename(columns={df.columns[0]: "date"}, inplace=True)
    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    return df.dropna(subset=["date"]).set_index("date").sort_index()


def lire_historique_colonne(nom_fichier, panier: str, colonne: str) -> pd.Series:
    """Série d'une colonne, indexée par date."""
    df = _lire_feuille(nom_fichier, panier)
    if colonne not in df.columns:
        return pd.Series(dtype=float)
    return df[colonne].dropna()


def derniers_mois(nom_fichier, panier: str, nb: int = 3) -> pd.DataFrame:
    """Les `nb` derniers mois connus du panier, pour repère à la saisie."""
    df = _lire_feuille(nom_fichier, panier)
    colonnes = [c for c in df.columns if not str(c).startswith("Unnamed")]
    return df[colonnes].tail(int(nb))


def dates_disponibles(nom_fichier, panier: str) -> list:
    """Mois déjà renseignés dans ce panier, du plus récent au plus ancien."""
    df = _lire_feuille(nom_fichier, panier)
    return list(df.index[::-1])


def lire_valeur(nom_fichier, panier: str, colonne: str, date_cible) -> float:
    """
    Valeur actuellement enregistrée pour un élément à une date donnée.

    Renvoie None si la colonne, le mois, ou la valeur elle-même sont absents.
    """
    df = _lire_feuille(nom_fichier, panier)
    if colonne not in df.columns:
        return None
    date_cible = pd.Timestamp(date_cible)
    ligne = df.loc[(df.index.year == date_cible.year) & (df.index.month == date_cible.month)]
    if ligne.empty:
        return None
    valeur = ligne.iloc[0][colonne]
    return None if pd.isna(valeur) else float(valeur)


def date_deja_presente(nom_fichier, panier: str, date_cible) -> bool:
    """Une ligne existe-t-elle déjà pour ce mois ? (cas d'une correction)"""
    date_cible = pd.Timestamp(date_cible)
    df = _lire_feuille(nom_fichier, panier)
    if df.empty:
        return False
    return bool(((df.index.year == date_cible.year) & (df.index.month == date_cible.month)).any())


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------


def valider_completude(valeurs: dict) -> str:
    """
    'vide', 'complet' ou 'partiel' — base de la règle du tout ou rien.

    Un panier à moitié rempli fausserait l'indice : l'IPC est une moyenne
    pondérée, une case vide n'est pas un zéro.
    """
    if not valeurs:
        return "vide"
    remplies = sum(1 for v in valeurs.values() if v is not None and not (isinstance(v, float) and pd.isna(v)))
    if remplies == 0:
        return "vide"
    if remplies == len(valeurs):
        return "complet"
    return "partiel"


# ---------------------------------------------------------------------------
# Écriture
# ---------------------------------------------------------------------------


def inserer_ligne_panier(nom_fichier, panier: str, date_cible, valeurs: dict, ecraser: bool = False) -> None:
    """
    Écrit une ligne de valeurs brutes pour `date_cible` dans `panier`.

    L'ordre des colonnes du classeur fait foi : les valeurs sont replacées
    par nom d'en-tête, jamais par position. Une colonne saisie mais absente
    de la feuille interrompt l'écriture plutôt que d'être écrite ailleurs.
    """
    chemin = _fichier(nom_fichier)
    date_cible = pd.Timestamp(date_cible)

    # Contrôler l'existence de la feuille AVANT toute lecture : sinon pandas
    # remonte son propre message, opaque pour l'utilisateur.
    if panier not in pd.ExcelFile(chemin).sheet_names:
        raise ValueError("Feuille « %s » absente de %s" % (panier, chemin))

    if date_deja_presente(chemin, panier, date_cible) and not ecraser:
        raise DateDejaPresente(
            "Une ligne existe déjà pour %04d-%02d dans « %s ». Demandez "
            "explicitement l'écrasement pour la corriger." % (date_cible.year, date_cible.month, panier)
        )

    classeur = load_workbook(chemin)
    if panier not in classeur.sheetnames:
        classeur.close()
        raise ValueError("Feuille « %s » absente de %s" % (panier, chemin))
    feuille = classeur[panier]

    # En-têtes : nom -> index de colonne
    entetes = {}
    for col in range(1, feuille.max_column + 1):
        nom = feuille.cell(row=1, column=col).value
        if nom is not None:
            entetes[str(nom).strip()] = col

    inconnues = [c for c in valeurs if str(c).strip() not in entetes]
    if inconnues:
        classeur.close()
        raise ValueError("Colonnes absentes de la feuille « %s » : %s" % (panier, ", ".join(sorted(inconnues))))

    # Ligne existante pour ce mois, sinon première ligne libre
    ligne_cible = None
    for rang in range(2, feuille.max_row + 1):
        cellule = feuille.cell(row=rang, column=1).value
        if cellule is None:
            continue
        try:
            date_ligne = pd.Timestamp(cellule)
        except (ValueError, TypeError):
            continue
        if date_ligne.year == date_cible.year and date_ligne.month == date_cible.month:
            ligne_cible = rang
            break

    if ligne_cible is None:
        ligne_cible = feuille.max_row + 1
        feuille.cell(row=ligne_cible, column=1, value=date_cible.to_pydatetime())

    for colonne, valeur in valeurs.items():
        if valeur is None or (isinstance(valeur, float) and pd.isna(valeur)):
            continue
        feuille.cell(row=ligne_cible, column=entetes[str(colonne).strip()], value=float(valeur))

    classeur.save(chemin)
    classeur.close()


def modifier_valeur(nom_fichier, panier: str, colonne: str, date_cible, nouvelle_valeur: float) -> None:
    """
    Corrige une valeur déjà enregistrée, sans toucher aux autres colonnes de
    la ligne. S'appuie sur inserer_ligne_panier, qui n'écrit jamais que les
    colonnes transmises : mêmes garde-fous d'écriture qu'à l'insertion.
    """
    inserer_ligne_panier(nom_fichier, panier, date_cible, {colonne: nouvelle_valeur}, ecraser=True)


def recalculer_apres_insertion(nom_fichier=None) -> None:
    """
    Relance la chaîne de calculs. Aucune logique d'indicateur ici : tout est
    dans calculator.pipeline_global().
    """
    pipeline_global(_fichier(nom_fichier))
