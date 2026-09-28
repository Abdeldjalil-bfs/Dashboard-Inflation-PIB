"""
Lecture des fichiers PIB de l'ONS (PIB_TR_S.xlsx = offre, PIB_TR_D.xlsx =
demande) et reformatage vers le schéma canonique de backend.pib.calculator.

Structure des fichiers ONS (vérifiée par réconciliation numérique, voir
docs/AUDIT_PIB.md) :
  - 'PIB_N' : valeurs aux prix courants ; 'PIB_R' : volumes aux prix de
    l'année précédente chaînés. Même disposition de colonnes, en-tête en
    ligne 4, données à partir de la ligne 5.
  - Offre : PIB = somme des branches + impôts nets sur les produits en
    NOMINAL (vérifié à l'arrondi près). En VOLUME, la somme des branches
    s'écarte du PIB publié (chaînage non additif) : le PIB réel est donc lu
    dans la colonne totale de l'ONS, jamais recalculé par somme.
  - Demande : PIB = C + G + FBCF + ΔS + X − M en nominal (vérifié exact).

Toute la correspondance libellés ONS -> clés canoniques vit dans
config/pib_config.json. Si l'ONS change la disposition de ses fichiers,
seuls ce module et ce fichier de configuration sont à ajuster.

Les saisies manuelles (journal config.settings.FICHIER_PIB_SAISIES) sont
appliquées ici, à la lecture : les classeurs ONS ne sont jamais réécrits.
"""

import json
import os
import re
import unicodedata

import openpyxl
import pandas as pd

# Colonnes du journal des saisies, dans l'ordre d'écriture.
COLONNES_JOURNAL = ["horodatage", "utilisateur", "fichier", "feuille", "date", "colonne", "action", "valeur"]


def config_pib() -> dict:
    from config.settings import PIB_CONFIG_PATH

    with open(PIB_CONFIG_PATH, "r", encoding="utf-8") as flux:
        return json.load(flux)


def normaliser(texte) -> str:
    """Insensible aux accents, à la casse et aux espaces multiples ou en trop :
    les intitulés ONS ont des espaces irréguliers ('Impôts sur les  importations')."""
    if texte is None:
        return ""
    texte = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", texte).strip().lower()


# ---------------------------------------------------------------------------
# Lecture brute d'un onglet
# ---------------------------------------------------------------------------


def lire_feuille_ons(fichier: str, feuille: str) -> pd.DataFrame:
    """
    Onglet ONS en DataFrame indexé par date de trimestre, une colonne par
    intitulé (espaces de bord retirés). En cas d'intitulé en double, la
    première colonne l'emporte (les totaux précèdent le sous-détail).
    """
    ligne_entete = config_pib()["feuilles"]["ligne_entete"]
    wb = openpyxl.load_workbook(fichier, read_only=True, data_only=True)
    try:
        ws = wb[feuille]
        lignes = list(ws.iter_rows(min_row=ligne_entete, values_only=True))
    finally:
        wb.close()

    entete = lignes[0]
    colonnes, vues = {}, set()
    for i, libelle in enumerate(entete):
        if i == 0 or libelle is None:
            continue
        cle = normaliser(libelle)
        if cle and cle not in vues:
            vues.add(cle)
            colonnes[i] = str(libelle).strip()

    donnees = {}
    for ligne in lignes[1:]:
        if ligne[0] is None:
            continue
        try:
            date = pd.Timestamp(ligne[0])
        except (TypeError, ValueError):
            continue
        donnees[date] = {nom: ligne[i] for i, nom in colonnes.items()}

    df = pd.DataFrame.from_dict(donnees, orient="index", columns=list(colonnes.values()))
    df = df.apply(pd.to_numeric, errors="coerce").sort_index()
    df.index.name = "date"
    return df


def titre_feuille(fichier: str, feuille: str) -> str:
    """Texte des lignes au-dessus de l'en-tête (titre ONS de l'onglet)."""
    ligne_entete = config_pib()["feuilles"]["ligne_entete"]
    wb = openpyxl.load_workbook(fichier, read_only=True, data_only=True)
    try:
        ws = wb[feuille]
        textes = [
            str(v)
            for ligne in ws.iter_rows(min_row=1, max_row=ligne_entete - 1, values_only=True)
            for v in ligne
            if isinstance(v, str)
        ]
    finally:
        wb.close()
    return " ".join(textes)


def colonne(df: pd.DataFrame, libelle: str) -> pd.Series:
    """Colonne de `df` dont l'intitulé normalisé égale celui de `libelle`."""
    cible = normaliser(libelle)
    for nom in df.columns:
        if normaliser(nom) == cible:
            return df[nom]
    raise KeyError(libelle)


def _verifier_colonnes(df, libelles, fichier, feuille):
    presentes = {normaliser(c) for c in df.columns}
    manquants = [l for l in libelles if normaliser(l) not in presentes]
    if manquants:
        raise ValueError(
            "Colonnes attendues introuvables dans '%s' (%s) : %s — la disposition du "
            "fichier ONS a peut-être changé, voir config/pib_config.json" % (fichier, feuille, manquants)
        )


# ---------------------------------------------------------------------------
# Journal des saisies
# ---------------------------------------------------------------------------


def lire_journal(journal=None) -> pd.DataFrame:
    """Journal des saisies PIB (vide si le fichier n'existe pas)."""
    if journal is None:
        from config.settings import FICHIER_PIB_SAISIES

        journal = FICHIER_PIB_SAISIES
    if not os.path.exists(str(journal)):
        return pd.DataFrame(columns=COLONNES_JOURNAL)
    df = pd.read_excel(str(journal))
    for nom in COLONNES_JOURNAL:
        if nom not in df.columns:
            df[nom] = None
    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    return df[COLONNES_JOURNAL].sort_values("horodatage", kind="stable")


def appliquer_journal(df: pd.DataFrame, fichier_cle: str, feuille: str, journal=None) -> pd.DataFrame:
    """
    Rejoue sur `df` les saisies du journal visant (`fichier_cle`, `feuille`),
    dans l'ordre chronologique. Insertion et modification posent la valeur
    (une nouvelle date crée la ligne) ; suppression la remplace par NaN.
    """
    entrees = lire_journal(journal)
    entrees = entrees[(entrees["fichier"] == fichier_cle) & (entrees["feuille"] == feuille)]
    if entrees.empty:
        return df

    df = df.copy()
    noms = {normaliser(c): c for c in df.columns}
    for _i, entree in entrees.iterrows():
        nom = noms.get(normaliser(entree["colonne"]))
        if nom is None or pd.isna(entree["date"]):
            continue
        date = pd.Timestamp(entree["date"])
        if entree["action"] == "suppression":
            if date in df.index:
                df.loc[date, nom] = float("nan")
        else:
            df.loc[date, nom] = float(entree["valeur"])
    return df.sort_index()


def lire_feuille_corrigee(fichier: str, fichier_cle: str, feuille: str, journal=None) -> pd.DataFrame:
    """Onglet ONS après application du journal des saisies."""
    return appliquer_journal(lire_feuille_ons(fichier, feuille), fichier_cle, feuille, journal)


# ---------------------------------------------------------------------------
# Schéma canonique
# ---------------------------------------------------------------------------


def _definitions_offre(config) -> dict:
    """{clé canonique: [intitulés ONS à sommer]} — secteurs puis impôts nets."""
    definitions = {s["cle"]: s["colonnes_ons"] for s in config["secteurs_offre"]}
    impots = config["impots_nets_produits"]
    definitions[impots["cle"]] = impots["colonnes_ons"]
    return definitions


def lire_offre_ons(fichier: str, journal=None) -> pd.DataFrame:
    """
    DataFrame indexé par trimestre : '<secteur>_nominal' / '<secteur>_reel'
    pour chaque secteur et les impôts nets, plus 'PIB_nominal' / 'PIB_reel'
    (totaux publiés par l'ONS).
    """
    config = config_pib()
    feuilles = config["feuilles"]
    definitions = _definitions_offre(config)
    attendues = [l for libelles in definitions.values() for l in libelles] + [config["colonne_pib_offre"]]

    resultat = {}
    for suffixe, feuille in (("nominal", feuilles["nominal"]), ("reel", feuilles["reel"])):
        df = lire_feuille_corrigee(fichier, "offre", feuille, journal)
        _verifier_colonnes(df, attendues, fichier, feuille)
        for cle, libelles in definitions.items():
            # min_count=1 : un trimestre non publié reste NaN, pas 0.
            resultat[f"{cle}_{suffixe}"] = pd.concat([colonne(df, l) for l in libelles], axis=1).sum(
                axis=1, min_count=1
            )
        resultat[f"PIB_{suffixe}"] = colonne(df, config["colonne_pib_offre"])

    return pd.DataFrame(resultat).sort_index()


def lire_demande_ons(fichier: str, journal=None) -> pd.DataFrame:
    """
    DataFrame indexé par trimestre : '<poste>_nominal' / '<poste>_reel' pour
    chaque poste de la demande, plus 'PIB_nominal' / 'PIB_reel'.
    """
    config = config_pib()
    feuilles = config["feuilles"]
    postes = config["postes_demande"]
    attendues = [p["colonne_ons"] for p in postes] + [config["colonne_pib_demande"]]

    resultat = {}
    for suffixe, feuille in (("nominal", feuilles["nominal"]), ("reel", feuilles["reel"])):
        df = lire_feuille_corrigee(fichier, "demande", feuille, journal)
        _verifier_colonnes(df, attendues, fichier, feuille)
        for poste in postes:
            resultat[f"{poste['cle']}_{suffixe}"] = colonne(df, poste["colonne_ons"])
        resultat[f"PIB_{suffixe}"] = colonne(df, config["colonne_pib_demande"])

    return pd.DataFrame(resultat).sort_index()


def colonnes_saisissables(fichier_cle: str) -> list:
    """Intitulés ONS lus par le tableau de bord pour ce fichier ('offre' / 'demande')."""
    config = config_pib()
    if fichier_cle == "offre":
        libelles = [l for ls in _definitions_offre(config).values() for l in ls]
        return libelles + [config["colonne_pib_offre"]]
    return [p["colonne_ons"] for p in config["postes_demande"]] + [config["colonne_pib_demande"]]


# ---------------------------------------------------------------------------
# Nature des volumes : prix constants ou volumes chaînés
# ---------------------------------------------------------------------------


def detecter_nature_volumes(fichier: str, df_offre: pd.DataFrame = None) -> dict:
    """
    Détermine si l'onglet réel porte des prix constants d'une année de base
    (additifs) ou des volumes chaînés (non additifs), sans le supposer :

      1. le titre ONS de l'onglet réel mentionne-t-il le chaînage ?
      2. la somme des branches en volume reconstitue-t-elle le PIB réel
         publié ? Un écart moyen supérieur au seuil configuré signe un
         chaînage.

    Les deux indices doivent concorder ; en cas de désaccord, le test
    numérique l'emporte (c'est lui qui conditionne l'additivité) et le
    désaccord est signalé dans `raison`.

    Retour : {"methode": "chainage" | "prix_constants", "raison", "titre",
              "ecart_additivite_moyen_pct", "indice_titre", "indice_numerique"}
    """
    config = config_pib()
    regles = config["methode_contributions"]
    feuille_reelle = config["feuilles"]["reel"]

    try:
        titre = titre_feuille(fichier, feuille_reelle)
    except Exception:
        titre = ""
    titre_norm = normaliser(titre)
    indice_titre = any(normaliser(m) in titre_norm for m in regles["mots_cles_chainage"])

    if df_offre is None:
        df_offre = lire_offre_ons(fichier)
    cles = [s["cle"] for s in config["secteurs_offre"]] + [config["impots_nets_produits"]["cle"]]
    somme = df_offre[[f"{c}_reel" for c in cles]].sum(axis=1, min_count=len(cles))
    ecart = ((somme - df_offre["PIB_reel"]) / df_offre["PIB_reel"] * 100).abs().dropna()
    ecart_moyen = float(ecart.mean()) if not ecart.empty else 0.0
    indice_numerique = ecart_moyen > regles["seuil_non_additivite_pct"]

    forcee = regles.get("choix", "auto")
    if forcee in ("chainage", "prix_constants"):
        methode = forcee
        raison = "Méthode imposée par config/pib_config.json."
    else:
        methode = "chainage" if indice_numerique else "prix_constants"
        raison = "Somme des branches en volume vs PIB réel publié : écart moyen de %.2f %% (seuil %.1f %%). " % (
            ecart_moyen,
            regles["seuil_non_additivite_pct"],
        )
        raison += (
            "Le titre de l'onglet réel mentionne le chaînage."
            if indice_titre
            else "Le titre de l'onglet réel ne mentionne pas le chaînage."
        )
        if indice_titre != indice_numerique:
            raison += " Indices discordants : le test numérique l'emporte."

    return {
        "methode": methode,
        "raison": raison,
        "titre": titre.strip(),
        "ecart_additivite_moyen_pct": round(ecart_moyen, 3),
        "indice_titre": indice_titre,
        "indice_numerique": indice_numerique,
    }
