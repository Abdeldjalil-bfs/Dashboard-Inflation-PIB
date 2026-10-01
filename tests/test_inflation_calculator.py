"""
Calculs Inflation sur un classeur factice : IPC pondéré, inflation MoM et
YoY, contributions (leur somme égale l'inflation du panier) et Core /
Non-Core. Les fonctions de calcul écrivent dans le classeur qu'on leur
passe : on travaille donc sur une copie temporaire, jamais sur les données.
"""

import json

import numpy as np
import pandas as pd
import pytest

from backend.inflation import calculator as C
from config.settings import WEIGHTS_PATH

POIDS = json.load(open(WEIGHTS_PATH, encoding="utf-8"))
DATES = pd.date_range("2020-01-01", periods=30, freq="MS")
DEBUT, FIN = "2020-01", "2022-06"


def _indices(colonnes, graine):
    rng = np.random.default_rng(graine)
    return pd.DataFrame({c: 100 * np.cumprod(1 + rng.normal(0.004, 0.01, len(DATES))) for c in colonnes}, index=DATES)


@pytest.fixture
def classeur(tmp_path):
    chemin = tmp_path / "calculs.xlsx"
    with pd.ExcelWriter(chemin) as xw:
        for rang, feuille in enumerate(("Grand_Alger", "core", "Produits_agricoles_frais")):
            _indices(list(POIDS[feuille]), rang).rename_axis("date").reset_index().to_excel(
                xw, sheet_name=feuille, index=False
            )
        cat = _indices(list(POIDS["categories"]), 9)
        cat.rename_axis("date").reset_index().to_excel(xw, sheet_name="categories", index=False)
    return str(chemin)


def _feuille(chemin, nom):
    return pd.read_excel(chemin, sheet_name=nom).set_index("date")


def test_ipc_pondere(classeur):
    C.calculer_ipc(classeur, "Grand_Alger", DEBUT, FIN)
    df = _feuille(classeur, "Grand_Alger")
    poids = pd.Series(POIDS["Grand_Alger"])
    attendu = (df[poids.index] * poids).sum(axis=1) / poids.sum()
    assert (df["IPC (%)"] - attendu).abs().max() < 1e-6


def test_inflation_mom_et_yoy(classeur):
    C.calculer_ipc(classeur, "Grand_Alger", DEBUT, FIN)
    C.calculer_inflation_mom(classeur, "Grand_Alger", DEBUT, FIN)
    C.calculer_inflation_yoy(classeur, "Grand_Alger", DEBUT, FIN)
    df = _feuille(classeur, "Grand_Alger")
    ipc = df["IPC (%)"]
    assert df["Inflation (%, mom)"].iloc[5] == pytest.approx((ipc.iloc[5] / ipc.iloc[4] - 1) * 100, abs=0.01)
    assert df["Inflation (%, yoy)"].iloc[20] == pytest.approx((ipc.iloc[20] / ipc.iloc[8] - 1) * 100, abs=0.01)
    assert df["Inflation (%, yoy)"].iloc[:12].isna().all()


@pytest.mark.parametrize("mode", ["mom", "yoy"])
def test_somme_des_contributions_egale_l_inflation(classeur, mode):
    for etape in (
        C.calculer_ipc,
        C.calculer_inflation_elements_mom,
        C.calculer_inflation_mom,
        C.calculer_inflation_elements_yoy,
        C.calculer_inflation_yoy,
        C.calculer_contributions_pp_mom,
        C.calculer_contributions_pp_yoy,
    ):
        etape(classeur, "Grand_Alger", DEBUT, FIN)
    for date in DATES[13:]:
        coherent, ecart = C.verifier_coherence_contributions(classeur, "Grand_Alger", date, mode, tolerance=0.05)
        assert coherent, (date, ecart)


def test_core_non_core(classeur):
    C.calculer_ipc(classeur, "categories", DEBUT, FIN)
    C.calculer_ipc_core_noncore(classeur, "core", "Produits_agricoles_frais", DEBUT, FIN)
    core = _feuille(classeur, "core")
    poids = pd.Series(POIDS["core"])
    attendu = (core[poids.index] * poids).sum(axis=1) / poids.sum()
    assert (core["IPC Core (%)"] - attendu).abs().max() < 1e-6
    frais = _feuille(classeur, "Produits_agricoles_frais")
    poids_f = pd.Series(POIDS["Produits_agricoles_frais"])
    attendu_f = (frais[poids_f.index] * poids_f).sum(axis=1) / poids_f.sum()
    assert (frais["IPC Non Core (%)"] - attendu_f).abs().max() < 1e-6


def _fichier_complementaire(tmp_path, dates, valeurs):
    """Onglet 'IPC_Catégories' minimal, au format large attendu par
    _parser_bloc_large : une ligne 'Poids' puis une colonne par mois."""
    from openpyxl import Workbook

    chemin = tmp_path / "complementaire.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "IPC_Catégories"
    ws.append(["Libellé", "Poids"] + list(dates))
    ws.append(["Produits agricoles frais", 169.18] + list(valeurs))
    wb.save(chemin)
    wb.close()
    return str(chemin)


def test_completer_ipc_non_core_depuis_complementaire(classeur, tmp_path):
    """
    Cas réel : un sous-produit du panier agricole frais manque sur les
    derniers mois dans le fichier source, 'IPC Non Core (%)' reste alors
    vide (la moyenne pondérée exige tous les postes) — on reprend l'agrégat
    déjà publié dans le fichier complémentaire pour ces mois précis, sans
    jamais écraser une case déjà renseignée par la somme pondérée.
    """
    from openpyxl import load_workbook

    C.calculer_ipc_core_noncore(classeur, "core", "Produits_agricoles_frais", DEBUT, FIN)

    wb = load_workbook(classeur)
    ws = wb["Produits_agricoles_frais"]
    col_ipc = next(
        c for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value == "IPC Non Core (%)"
    )
    lignes_manquantes = [ws.max_row - 1, ws.max_row]
    dates_manquantes = [ws.cell(row=r, column=1).value for r in lignes_manquantes]
    for r in lignes_manquantes:
        ws.cell(row=r, column=col_ipc).value = None  # cell(..., value=None) ne vide pas la cellule : il faut l'attribut
    wb.save(classeur)
    wb.close()

    valeurs_agregat = [123.45, 126.78]
    complementaire = _fichier_complementaire(tmp_path, dates_manquantes, valeurs_agregat)

    assert C.completer_ipc_non_core_depuis_complementaire(classeur, "Produits_agricoles_frais", complementaire) is True

    frais = _feuille(classeur, "Produits_agricoles_frais")
    for date, valeur in zip(dates_manquantes, valeurs_agregat):
        assert frais.loc[pd.Timestamp(date), "IPC Non Core (%)"] == pytest.approx(valeur)

    # Idempotent, et une case déjà renseignée (toutes les autres) n'est
    # jamais écrasée : un second passage ne comble donc plus rien.
    assert (
        C.completer_ipc_non_core_depuis_complementaire(classeur, "Produits_agricoles_frais", complementaire) is False
    )
